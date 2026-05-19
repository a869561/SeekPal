"""Smoke tests del conftest: verifica que las fixtures generan ficheros válidos."""

from pathlib import Path


def test_text_fixture_creates_utf8_file(tmp_text_file: Path):
    assert tmp_text_file.exists()
    assert tmp_text_file.suffix == ".txt"
    content = tmp_text_file.read_text(encoding="utf-8")
    assert "SeekPal" in content
    assert "RAG" in content


def test_md_fixture_has_tildes(tmp_md_file: Path):
    assert tmp_md_file.exists()
    content = tmp_md_file.read_text(encoding="utf-8")
    assert "año" in content
    assert "Título" in content


def test_pdf_fixture_has_two_pages(tmp_pdf_file: Path):
    import fitz

    assert tmp_pdf_file.exists()
    doc = fitz.open(str(tmp_pdf_file))
    try:
        assert doc.page_count == 2
    finally:
        doc.close()


def test_docx_fixture_loads(tmp_docx_file: Path):
    from docx import Document

    assert tmp_docx_file.exists()
    doc = Document(str(tmp_docx_file))
    paras = [p.text for p in doc.paragraphs]
    assert any("SeekPal" in p for p in paras)


def test_pptx_fixture_has_two_slides(tmp_pptx_file: Path):
    from pptx import Presentation

    assert tmp_pptx_file.exists()
    prs = Presentation(str(tmp_pptx_file))
    assert len(prs.slides) == 2


def test_odt_fixture_loads(tmp_odt_file: Path):
    from odf.opendocument import load

    assert tmp_odt_file.exists()
    doc = load(str(tmp_odt_file))
    assert doc is not None
